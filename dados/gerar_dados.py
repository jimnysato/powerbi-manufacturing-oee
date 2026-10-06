"""Gera os dados fictícios da "Fábrica Exemplo" (injeção plástica) para a demo B do portfólio.

Imita o export de um sistema de apontamento de produção:
- Apontamentos.xlsx: 1 linha por injetora x turno x dia, com até 5 campos de parada na mesma linha,
  data em texto dd/mm/aaaa e turno da noite cruzando a meia-noite.
- Cadastros.xlsx: Injetoras, Produtos, Motivos e Turnos (com nomes em PT e EN).
- ~3% dos apontamentos têm erros de propósito (motivo vazio, quantidade vazia, hora de fim vazia),
  para alimentar a página de qualidade dos apontamentos.

Dados 100% inventados. Semente fixa: rodar de novo gera exatamente os mesmos arquivos.

Uso: python gerar_dados.py
"""
import os
import random
from datetime import date, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font

random.seed(42)
PASTA = os.path.dirname(os.path.abspath(__file__))
INICIO, FIM = date(2026, 4, 1), date(2026, 9, 30)
MIN_TURNO = 480

TURNOS = [
    ("T1", "1º turno", "1st shift", "06:00", "14:00"),
    ("T2", "2º turno", "2nd shift", "14:00", "22:00"),
    ("T3", "3º turno", "3rd shift", "22:00", "06:00"),
]

# código, toneladas, ano, fator de desempenho, taxa de refugo, paradas por turno (média)
INJETORAS = [
    ("INJ-01", 450, 2022, 0.96, 0.015, 0.9),
    ("INJ-02", 450, 2021, 0.94, 0.020, 1.0),
    ("INJ-03", 300, 2019, 0.93, 0.025, 1.1),
    ("INJ-04", 300, 2018, 0.91, 0.030, 1.3),
    ("INJ-05", 250, 2016, 0.90, 0.035, 1.4),
    ("INJ-06", 250, 2009, 0.82, 0.060, 2.2),
    ("INJ-07", 150, 2020, 0.95, 0.020, 1.0),
    ("INJ-08", 150, 2014, 0.88, 0.040, 1.6),
]

# código, descrição PT, descrição EN, família PT, família EN, ciclo (s), cavidades, peso (g), tonelagem mínima
PRODUTOS = [
    ("P-1001", "Tampa rosca 28 mm", "Screw cap 28 mm", "Tampas", "Caps", 8, 32, 2.1, 150),
    ("P-1002", "Tampa rosca 38 mm", "Screw cap 38 mm", "Tampas", "Caps", 9, 24, 3.4, 150),
    ("P-1003", "Tampa flip-top 24 mm", "Flip-top cap 24 mm", "Tampas", "Caps", 11, 16, 2.8, 150),
    ("P-1004", "Tampa de pote 500 ml", "Lid for 500 ml jar", "Tampas", "Caps", 10, 8, 9.5, 150),
    ("P-1005", "Tampa de balde 10 L", "Lid for 10 L bucket", "Tampas", "Caps", 28, 1, 140.0, 300),
    ("P-2001", "Pote redondo 250 ml", "Round jar 250 ml", "Potes", "Jars", 12, 8, 14.0, 150),
    ("P-2002", "Pote redondo 500 ml", "Round jar 500 ml", "Potes", "Jars", 14, 4, 22.0, 250),
    ("P-2003", "Pote quadrado 1 L", "Square container 1 L", "Potes", "Jars", 16, 4, 38.0, 250),
    ("P-2004", "Pote hermético 2 L", "Airtight container 2 L", "Potes", "Jars", 20, 2, 75.0, 300),
    ("P-2005", "Copo descartável 300 ml", "Disposable cup 300 ml", "Potes", "Jars", 6, 16, 4.5, 150),
    ("P-3001", "Balde 10 L", "Bucket 10 L", "Baldes", "Buckets", 32, 1, 420.0, 450),
    ("P-3002", "Balde 20 L", "Bucket 20 L", "Baldes", "Buckets", 40, 1, 780.0, 450),
    ("P-3003", "Bacia 15 L", "Basin 15 L", "Baldes", "Buckets", 35, 1, 510.0, 450),
    ("P-3004", "Regador 5 L", "Watering can 5 L", "Baldes", "Buckets", 30, 1, 310.0, 300),
    ("P-4001", "Caixa organizadora 5 L", "Storage box 5 L", "Organizadores", "Storage", 26, 1, 260.0, 300),
    ("P-4002", "Caixa organizadora 12 L", "Storage box 12 L", "Organizadores", "Storage", 34, 1, 480.0, 450),
    ("P-4003", "Gaveteiro (gaveta)", "Drawer unit (drawer)", "Organizadores", "Storage", 30, 1, 350.0, 450),
    ("P-4004", "Cesto multiuso", "Multi-purpose basket", "Organizadores", "Storage", 22, 1, 190.0, 300),
    ("P-4005", "Porta-talheres", "Cutlery tray", "Organizadores", "Storage", 18, 2, 95.0, 250),
    ("P-5001", "Cabide adulto", "Adult hanger", "Utilidades", "Household", 10, 4, 28.0, 150),
    ("P-5002", "Cabide infantil", "Kids hanger", "Utilidades", "Household", 9, 4, 18.0, 150),
    ("P-5003", "Prendedor de roupa", "Clothes peg", "Utilidades", "Household", 7, 24, 3.0, 150),
    ("P-5004", "Escorredor de louça", "Dish drainer", "Utilidades", "Household", 28, 1, 330.0, 300),
    ("P-5005", "Pá de lixo", "Dustpan", "Utilidades", "Household", 20, 2, 120.0, 250),
    ("P-5006", "Vaso de planta 3 L", "Plant pot 3 L", "Utilidades", "Household", 18, 2, 110.0, 250),
    ("P-6001", "Conector de mangueira", "Hose connector", "Técnicos", "Technical", 12, 8, 6.0, 150),
    ("P-6002", "Caixa de passagem elétrica", "Electrical junction box", "Técnicos", "Technical", 16, 4, 45.0, 250),
    ("P-6003", "Engrenagem nylon 40 mm", "Nylon gear 40 mm", "Técnicos", "Technical", 14, 8, 12.0, 150),
    ("P-6004", "Carcaça de ventilador", "Fan housing", "Técnicos", "Technical", 38, 1, 640.0, 450),
    ("P-6005", "Suporte de prateleira", "Shelf bracket", "Técnicos", "Technical", 15, 4, 32.0, 250),
]

# código, motivo PT, motivo EN, tipo PT, tipo EN, peso de sorteio, minutos (mín, máx)
MOTIVOS = [
    ("M01", "Setup / troca de molde", "Setup / mold change", "Planejada", "Planned", 0, (40, 120)),
    ("M02", "Falta de matéria-prima", "Raw material shortage", "Logística", "Logistics", 22, (20, 150)),
    ("M03", "Manutenção corretiva", "Corrective maintenance", "Manutenção", "Maintenance", 18, (30, 240)),
    ("M04", "Ajuste de molde", "Mold adjustment", "Processo", "Process", 20, (15, 60)),
    ("M05", "Problema de qualidade", "Quality issue", "Qualidade", "Quality", 12, (10, 50)),
    ("M06", "Falta de operador", "No operator available", "Pessoas", "People", 8, (30, 120)),
    ("M07", "Troca de cor", "Color change", "Planejada", "Planned", 12, (20, 60)),
    ("M08", "Outros", "Other", "Outros", "Other", 8, (5, 30)),
]
MOTIVOS_SORTEIO = [m for m in MOTIVOS if m[5] > 0]

MATRICULAS = [f"MAT-{n}" for n in range(101, 131)]


def poisson(lam):
    """Sorteio de Poisson simples (algoritmo de Knuth)."""
    limite, k, p = 2.718281828 ** -lam, 0, 1.0
    while True:
        p *= random.random()
        if p <= limite:
            return k
        k += 1


def produtos_da_injetora(ton):
    return [p for p in PRODUTOS if p[8] <= ton and p[8] >= ton - 200] or [p for p in PRODUTOS if p[8] <= ton]


def gerar_apontamentos():
    linhas = []
    op_seq = 26000
    estado = {}  # injetora -> [produto, OP, turnos restantes]
    dia = INICIO
    while dia <= FIM:
        dom, sab = dia.weekday() == 6, dia.weekday() == 5
        for inj, ton, _ano, desemp, refugo_base, lam in INJETORAS:
            if dom:
                continue
            for cod_t, _, _, h_ini, h_fim in TURNOS:
                if sab and cod_t == "T3":
                    continue
                if random.random() < 0.03:  # turno sem programação (máquina parada o turno todo, sem apontamento)
                    continue
                prod_atual = estado.get(inj)
                nova_op = prod_atual is None or prod_atual[2] <= 0
                if nova_op:
                    op_seq += 1
                    prod = random.choice(produtos_da_injetora(ton))
                    estado[inj] = [prod, f"OP-{op_seq}", random.randint(4, 15)]
                prod, op, _ = estado[inj]
                estado[inj][2] -= 1
                _c, _d, _de, _f, _fe, ciclo, cav, _peso, _t = prod

                paradas = []
                if nova_op:
                    m = MOTIVOS[0]
                    paradas.append([m[1], random.randint(*m[6])])
                fator_noite = 1.25 if cod_t == "T3" else 1.0
                fator_tendencia = 1.0 - 0.15 * ((dia - INICIO).days / (FIM - INICIO).days) if inj != "INJ-06" else 1.0
                for _ in range(poisson(lam * fator_noite * fator_tendencia)):
                    if len(paradas) >= 5:
                        break
                    m = random.choices(MOTIVOS_SORTEIO, weights=[x[5] for x in MOTIVOS_SORTEIO])[0]
                    if inj == "INJ-06" and random.random() < 0.35:
                        m = MOTIVOS[2]  # máquina antiga: mais manutenção corretiva
                    paradas.append([m[1], random.randint(*m[6])])
                total = sum(p[1] for p in paradas)
                if total > 360:
                    escala = 360 / total
                    for p in paradas:
                        p[1] = max(5, int(p[1] * escala))
                    total = sum(p[1] for p in paradas)

                prevista = int(MIN_TURNO * 60 / ciclo * cav)
                operando = MIN_TURNO - total
                produzida = int(operando * 60 / ciclo * cav * min(1.0, random.gauss(desemp, 0.03)))
                taxa_ref = max(0.0, random.gauss(refugo_base * (1.2 if cod_t == "T3" else 1.0), refugo_base * 0.3))
                refugo = int(produzida * taxa_ref)

                linha = {
                    "Data": dia.strftime("%d/%m/%Y"), "Turno": cod_t, "Máquina": inj, "OP": op,
                    "Cód. Produto": prod[0], "Matrícula": random.choice(MATRICULAS),
                    "Hora Início": h_ini, "Hora Fim": h_fim,
                    "Qtd Prevista": prevista, "Qtd Produzida": produzida, "Qtd Refugo": refugo,
                }
                for i in range(5):
                    linha[f"Motivo Parada {i + 1}"] = paradas[i][0] if i < len(paradas) else None
                    linha[f"Min Parada {i + 1}"] = paradas[i][1] if i < len(paradas) else None

                # erros de propósito (~3%)
                r = random.random()
                if r < 0.012 and paradas:
                    linha["Motivo Parada 1"] = None
                elif r < 0.022:
                    linha["Qtd Produzida"] = None
                elif r < 0.030:
                    linha["Hora Fim"] = None
                linhas.append(linha)
        dia += timedelta(days=1)
    for i, l in enumerate(linhas, start=1):
        l["ID Apontamento"] = 100000 + i
    return linhas


def salvar(caminho, abas):
    wb = Workbook()
    wb.remove(wb.active)
    for nome, cab, linhas in abas:
        ws = wb.create_sheet(nome)
        ws.append(cab)
        for c in ws[1]:
            c.font = Font(bold=True)
        for l in linhas:
            ws.append(l)
    wb.save(caminho)


def main():
    aps = gerar_apontamentos()
    cab = ["ID Apontamento", "Data", "Turno", "Máquina", "OP", "Cód. Produto", "Matrícula", "Hora Início", "Hora Fim",
           "Qtd Prevista", "Qtd Produzida", "Qtd Refugo"]
    for i in range(1, 6):
        cab += [f"Motivo Parada {i}", f"Min Parada {i}"]
    salvar(os.path.join(PASTA, "Apontamentos.xlsx"), [("Apontamentos", cab, [[a[c] for c in cab] for a in aps])])

    salvar(os.path.join(PASTA, "Cadastros.xlsx"), [
        ("Injetoras", ["Máquina", "Tonelagem", "Ano Fabricação"], [[i[0], i[1], i[2]] for i in INJETORAS]),
        ("Produtos", ["Cód. Produto", "Descrição", "Description", "Família", "Family", "Ciclo (s)", "Cavidades", "Peso (g)"],
         [[p[0], p[1], p[2], p[3], p[4], p[5], p[6], p[7]] for p in PRODUTOS]),
        ("Motivos", ["Cód. Motivo", "Motivo", "Reason", "Tipo", "Type"], [[m[0], m[1], m[2], m[3], m[4]] for m in MOTIVOS]),
        ("Turnos", ["Turno", "Nome", "Name", "Início", "Fim"], [list(t) for t in TURNOS]),
    ])

    n = len(aps)
    sem_motivo = sum(1 for a in aps if a["Motivo Parada 1"] is None and a["Min Parada 1"])
    sem_qtd = sum(1 for a in aps if a["Qtd Produzida"] is None)
    sem_fim = sum(1 for a in aps if a["Hora Fim"] is None)
    print(f"{n} apontamentos | sem motivo: {sem_motivo} | sem qtd: {sem_qtd} | sem hora fim: {sem_fim}")


if __name__ == "__main__":
    main()
